"""Document reading: turn an attachment file into raw text, regardless of format.

Native text (txt/pdf/docx) is extracted directly. Image-only PDFs and image
attachments fall back to Gemini vision OCR when available; otherwise they're
reported unreadable so the pipeline can escalate instead of guessing.
"""

from dataclasses import dataclass
from pathlib import Path

from ..llm import gemini

TEXT_EXTS = {".txt"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg"}


@dataclass(frozen=True)
class IngestResult:
    text: str
    method: str  # plain_text | pdf_native | docx | pdf_ocr_llm | image_ocr_llm | unreadable
    warning: str | None = None


def _read_pdf_native(path: Path) -> str:
    import pymupdf
    chunks = []
    with pymupdf.open(path) as doc:
        for page in doc:
            chunks.append(page.get_text())
            for table in page.find_tables().tables:
                for row in table.extract():
                    chunks.append(" | ".join(cell or "" for cell in row))
    return "\n".join(c for c in chunks if c).strip()


def _render_pdf_pages(path: Path) -> list[bytes]:
    import pymupdf
    with pymupdf.open(path) as doc:
        return [page.get_pixmap(dpi=200).tobytes("png") for page in doc]


def _read_docx(path: Path) -> str:
    import docx
    doc = docx.Document(path)
    chunks = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(chunks).strip()


def read_attachment(path: str, use_llm_ocr: bool = True) -> IngestResult:
    p = Path(path)
    if not p.exists():
        return IngestResult(text="", method="unreadable", warning="file not found")

    suffix = p.suffix.lower()

    if suffix in TEXT_EXTS:
        return IngestResult(text=p.read_text(errors="ignore"), method="plain_text")

    if suffix in DOCX_EXTS:
        try:
            text = _read_docx(p)
        except Exception as e:
            return IngestResult(text="", method="unreadable", warning=str(e))
        return (
            IngestResult(text=text, method="docx")
            if text else IngestResult(text="", method="unreadable", warning="empty document")
        )

    if suffix in PDF_EXTS:
        try:
            text = _read_pdf_native(p)
        except Exception as e:
            return IngestResult(text="", method="unreadable", warning=str(e))
        if text:
            return IngestResult(text=text, method="pdf_native")

        if use_llm_ocr and gemini.is_available():
            try:
                pages_text = [gemini.read_image_text(img, "image/png") for img in _render_pdf_pages(p)]
            except Exception as e:
                return IngestResult(text="", method="unreadable", warning=str(e))
            text = "\n".join(t for t in pages_text if t)
            if text:
                return IngestResult(text=text, method="pdf_ocr_llm")
        return IngestResult(text="", method="unreadable", warning="scanned PDF, no OCR available")

    if suffix in IMAGE_EXTS:
        if use_llm_ocr and gemini.is_available():
            text = gemini.read_image_text(p.read_bytes(), f"image/{suffix.lstrip('.')}")
            if text:
                return IngestResult(text=text, method="image_ocr_llm")
        return IngestResult(text="", method="unreadable", warning="image attachment, no OCR available")

    return IngestResult(text="", method="unreadable", warning=f"unsupported file type: {suffix}")
