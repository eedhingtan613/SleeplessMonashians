import pymupdf
import docx as docx_lib

from sdoc.core.ingest import read_attachment
from sdoc.llm import gemini


def test_plain_text(tmp_path):
    p = tmp_path / "si.txt"
    p.write_text("Shipper: ABC Co.")
    result = read_attachment(str(p))
    assert result.method == "plain_text"
    assert "ABC Co." in result.text


def test_native_pdf_text(tmp_path):
    p = tmp_path / "si.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Shipper: ABC Trading Co.")
    doc.save(p)
    result = read_attachment(str(p))
    assert result.method == "pdf_native"
    assert "ABC Trading Co." in result.text


def test_docx_paragraphs_and_tables(tmp_path):
    p = tmp_path / "si.docx"
    d = docx_lib.Document()
    d.add_paragraph("Shipper: ABC Trading Co.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Container Count"
    table.rows[0].cells[1].text = "3"
    d.save(p)
    result = read_attachment(str(p))
    assert result.method == "docx"
    assert "ABC Trading Co." in result.text
    assert "Container Count | 3" in result.text


def test_scanned_pdf_falls_back_to_llm_ocr(tmp_path, monkeypatch):
    p = tmp_path / "scanned.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    blank = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 100, 100), False)
    blank.set_rect(blank.irect, (255, 255, 255))
    page.insert_image(pymupdf.Rect(0, 0, 100, 100), pixmap=blank)
    doc.save(p)

    monkeypatch.setattr(gemini, "is_available", lambda: True)
    monkeypatch.setattr(gemini, "read_image_text", lambda img, mime: "Shipper: OCR Co.")

    result = read_attachment(str(p))
    assert result.method == "pdf_ocr_llm"
    assert result.text == "Shipper: OCR Co."


def test_scanned_pdf_without_ocr_is_unreadable(tmp_path, monkeypatch):
    p = tmp_path / "scanned.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    blank = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 100, 100), False)
    blank.set_rect(blank.irect, (255, 255, 255))
    page.insert_image(pymupdf.Rect(0, 0, 100, 100), pixmap=blank)
    doc.save(p)

    monkeypatch.setattr(gemini, "is_available", lambda: False)
    result = read_attachment(str(p))
    assert result.method == "unreadable"


def test_missing_file():
    result = read_attachment("/no/such/file.pdf")
    assert result.method == "unreadable"
